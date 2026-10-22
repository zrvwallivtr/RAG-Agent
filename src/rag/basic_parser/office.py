from pathlib import Path

from src.libs.constants import ParserResult
from src.logger import app_logger

app_log = app_logger(f"{__name__}.app")


def parse_docx(path: Path) -> ParserResult | None:
    """Extracts structural text elements line by line from document."""
    try:
        import docx

        doc = docx.Document(path)
        paragraphs = []

        for para in doc.paragraphs:
            clean_text = para.text.strip()
            if clean_text:
                paragraphs.append(clean_text)

        app_log.info(
            "Text content extracted from DOCX file '%s' line by line: Paragraphs=%d",
            path,
            len(paragraphs),
        )
        return "\n".join(paragraphs), "txt"

    except Exception as e:
        app_log.warning("Failed to read DOCX file '%s': %s", path, e)
        return


def parse_xlsx(path: Path) -> ParserResult | None:
    """Extracts text content from spreadsheet row by row."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        excel_text = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            excel_text.append(f"=== Sheet: {sheet_name} ===")

            for row in ws.iter_rows(values_only=True):
                # Only process lines that contain actual data elements
                if any(cell is not None for cell in row):
                    clean_row = [
                        str(cell).strip() if cell is not None else "" for cell in row
                    ]
                    excel_text.append(" | ".join(clean_row))

        app_log.info(
            "Text content extracted from XLSX file '%s' row by row: Sheets=%d",
            path,
            len(wb.sheetnames),
        )
        return "\n".join(excel_text), "txt"

    except Exception as e:
        app_log.warning("Failed to read XLSX file '%s': %s", path, e)
        return
